# -*- coding: utf-8 -*-
"""选项卡④ 地址坐标互转：单点查询 + 批量处理（地址 <-> 坐标）。

复用 core/geo_service 的多 provider 服务层 + core/utils 的度分秒解析。
"""
import csv
import os
import shutil
import subprocess
from datetime import datetime

from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGroupBox, QGridLayout, QLabel,
    QLineEdit, QComboBox, QPushButton, QRadioButton, QButtonGroup,
    QTextEdit, QTabWidget, QTableWidget, QTableWidgetItem, QFileDialog,
    QMessageBox, QProgressBar, QHeaderView, QCheckBox, QApplication,
)

from openpyxl import load_workbook, Workbook
from core import geo_service
from core.config import load_config
from core.utils import parse_coordinate_str, validate_coordinate
from core.coords import transform_between
from core.batch_pipeline import inspect_table, stream_address_conversion

CRS_LIST = ['WGS84', 'CGCS2000', 'GCJ-02', 'BD-09']
CSV_ENCODINGS = ['utf-8-sig', 'gbk', 'gb2312', 'utf-8']


def _preferred():
    return load_config().get('defaults', {}).get('preferred_provider', 'tianditu')


class SingleQueryWorker(QThread):
    result = Signal(dict)
    error = Signal(str)

    def __init__(self, mode, address, city, lon, lat, input_crs):
        super().__init__()
        self.mode = mode
        self.address = address
        self.city = city
        self.lon = lon
        self.lat = lat
        self.input_crs = input_crs

    def run(self):
        preferred = _preferred()
        if self.mode == 'addr2coord':
            r = geo_service.geocode_with_fallback(self.address, self.city, '', preferred)
            if r:
                self.result.emit({'type': 'addr', 'data': r})
            else:
                self.error.emit('地址解析失败，请检查网络或 API Key。')
        else:
            wlon, wlat = transform_between(self.lon, self.lat, self.input_crs, 'WGS84')
            addr = geo_service.reverse_full(wlon, wlat, preferred)
            if addr:
                self.result.emit({'type': 'coord', 'data': addr})
            else:
                self.error.emit('逆地址解析失败，请检查网络或 API Key。')


class BatchWorker(QThread):
    progress = Signal(int, int)
    log = Signal(str)
    finished_ok = Signal(dict)
    error = Signal(str)

    def __init__(self, mode, input_path, output_path, addr_col, lon_col, lat_col, crs_col, city, province):
        super().__init__()
        self.mode = mode
        self.input_path = input_path
        self.output_path = output_path
        self.addr_col = addr_col
        self.lon_col = lon_col
        self.lat_col = lat_col
        self.crs_col = crs_col
        self.city = city
        self.province = province
        self._cancelled = False

    def stop(self):
        self._cancelled = True

    def run(self):
        try:
            result = stream_address_conversion(
                self.input_path, self.output_path, self.mode, self.addr_col,
                self.lon_col, self.lat_col, self.crs_col, self.city, self.province,
                progress_cb=lambda done, total: self.progress.emit(done, total),
                log_cb=self.log.emit, cancel_cb=lambda: self._cancelled,
            )
            self.finished_ok.emit(result)
        except Exception as exc:
            self.error.emit(str(exc))


class TabAddrGeocode(QWidget):
    def __init__(self):
        super().__init__()
        self._init_ui()

    # ------------------------------------------------------------------ UI
    def _init_ui(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(16, 16, 16, 16)
        outer.setAlignment(Qt.AlignTop)
        self.sub_tabs = QTabWidget()
        outer.addWidget(self.sub_tabs)
        self.sub_tabs.addTab(self._build_single_page(), '单点查询')
        self.sub_tabs.addTab(self._build_batch_page(), '批量处理')

    # -------------------------------------------------------- 单点查询页
    def _build_single_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setAlignment(Qt.AlignTop)

        grp = QGroupBox('查询输入')
        g = QVBoxLayout()

        mode_row = QHBoxLayout()
        self.rb_addr2coord = QRadioButton('地址 → 坐标')
        self.rb_coord2addr = QRadioButton('坐标 → 地址')
        self.rb_addr2coord.setChecked(True)
        self.mode_group = QButtonGroup(self)
        self.mode_group.addButton(self.rb_addr2coord, 1)
        self.mode_group.addButton(self.rb_coord2addr, 2)
        mode_row.addWidget(self.rb_addr2coord)
        mode_row.addWidget(self.rb_coord2addr)
        mode_row.addStretch()
        g.addLayout(mode_row)

        addr_row = QHBoxLayout()
        addr_row.addWidget(QLabel('地址：'))
        self.addr_edit = QLineEdit()
        self.addr_edit.setPlaceholderText('如：民族大道 1 号')
        addr_row.addWidget(self.addr_edit, 1)
        addr_row.addWidget(QLabel('限定城市：'))
        self.city_edit = QLineEdit()
        self.city_edit.setPlaceholderText('如：南宁市')
        addr_row.addWidget(self.city_edit)
        g.addLayout(addr_row)

        coord_row = QHBoxLayout()
        coord_row.addWidget(QLabel('经度：'))
        self.lon_edit = QLineEdit()
        self.lon_edit.setPlaceholderText('108.3667 或 108°22′00″E')
        coord_row.addWidget(self.lon_edit)
        coord_row.addWidget(QLabel('纬度：'))
        self.lat_edit = QLineEdit()
        self.lat_edit.setPlaceholderText('22.8167')
        coord_row.addWidget(self.lat_edit)
        coord_row.addWidget(QLabel('输入坐标系：'))
        self.input_crs_combo = QComboBox()
        self.input_crs_combo.addItems(CRS_LIST)
        coord_row.addWidget(self.input_crs_combo)
        g.addLayout(coord_row)

        btn_row = QHBoxLayout()
        self.query_btn = QPushButton('查询')
        self.query_btn.clicked.connect(self._on_query)
        self.clear_btn = QPushButton('清空')
        self.clear_btn.clicked.connect(self._on_clear)
        btn_row.addWidget(self.query_btn)
        btn_row.addWidget(self.clear_btn)
        btn_row.addStretch()
        g.addLayout(btn_row)

        grp.setLayout(g)
        layout.addWidget(grp)

        res_grp = QGroupBox('转换结果')
        rg = QVBoxLayout()
        self.result_text = QTextEdit()
        self.result_text.setReadOnly(True)
        self.result_text.setMinimumHeight(180)
        rg.addWidget(self.result_text)
        copy_btn = QPushButton('复制结果')
        copy_btn.clicked.connect(self._copy_result)
        rg.addWidget(copy_btn, 0, Qt.AlignRight)
        res_grp.setLayout(rg)
        layout.addWidget(res_grp)

        self.rb_addr2coord.toggled.connect(self._toggle_input_mode)
        self._toggle_input_mode()
        return page

    def _toggle_input_mode(self):
        is_addr = self.rb_addr2coord.isChecked()
        self.addr_edit.setEnabled(is_addr)
        self.city_edit.setEnabled(is_addr)
        self.lon_edit.setEnabled(not is_addr)
        self.lat_edit.setEnabled(not is_addr)
        self.input_crs_combo.setEnabled(not is_addr)

    def _on_query(self):
        if self.rb_addr2coord.isChecked():
            address = self.addr_edit.text().strip()
            if not address:
                QMessageBox.warning(self, '提示', '请输入地址。')
                return
            city = self.city_edit.text().strip()
            self.result_text.setText('查询中…')
            self.query_btn.setEnabled(False)
            self.worker = SingleQueryWorker('addr2coord', address, city, 0, 0, '')
            self.worker.result.connect(self._on_single_result)
            self.worker.error.connect(self._on_single_error)
            self.worker.finished.connect(lambda: self.query_btn.setEnabled(True))
            self.worker.start()
        else:
            lon = parse_coordinate_str(self.lon_edit.text())
            lat = parse_coordinate_str(self.lat_edit.text())
            if lon is None or lat is None or not validate_coordinate(lon, lat):
                QMessageBox.warning(self, '提示', '经纬度格式不正确或超出范围。')
                return
            input_crs = self.input_crs_combo.currentText()
            self.result_text.setText('查询中…')
            self.query_btn.setEnabled(False)
            self.worker = SingleQueryWorker('coord2addr', '', '', lon, lat, input_crs)
            self.worker.result.connect(self._on_single_result)
            self.worker.error.connect(self._on_single_error)
            self.worker.finished.connect(lambda: self.query_btn.setEnabled(True))
            self.worker.start()

    def _on_single_result(self, payload):
        if payload['type'] == 'addr':
            d = payload['data']
            w = d['wgs84']; gcj = d['gcj02']; bd = d['bd09']
            self.result_text.setText(
                f'标准地址：{d["formatted"]}\n'
                f'来源：{d["provider"]}\n\n'
                f'WGS84：{w[0]:.6f}, {w[1]:.6f}\n'
                f'GCJ-02：{gcj[0]:.6f}, {gcj[1]:.6f}\n'
                f'BD-09：{bd[0]:.6f}, {bd[1]:.6f}'
            )
        else:
            a = payload['data']
            self.result_text.setText(
                f'标准地址：{a["formatted"]}\n'
                f'来源：{a.get("provider", "")}\n\n'
                f'省：{a["province"]}\n'
                f'市：{a["city"]}\n'
                f'区/县：{a["district"]}\n'
                f'乡镇/街道：{a.get("township", "")} {a["street"]}\n'
                f'门牌号：{a["number"]}'
            )

    def _on_single_error(self, msg):
        self.result_text.setText(f'错误：{msg}')

    def _on_clear(self):
        self.addr_edit.clear()
        self.lon_edit.clear()
        self.lat_edit.clear()
        self.result_text.clear()

    def _copy_result(self):
        text = self.result_text.toPlainText()
        if text:
            QApplication.clipboard().setText(text)
            QMessageBox.information(self, '提示', '已复制到剪贴板。')

    # -------------------------------------------------------- 批量处理页
    def _build_batch_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setAlignment(Qt.AlignTop)

        file_grp = QGroupBox('文件导入')
        fg = QHBoxLayout()
        self.file_edit = QLineEdit()
        self.file_edit.setReadOnly(True)
        self.file_edit.setPlaceholderText('选择 Excel（.xlsx/.xls）或 CSV（.csv）')
        fg.addWidget(self.file_edit, 1)
        browse_btn = QPushButton('浏览…')
        browse_btn.clicked.connect(self._browse)
        fg.addWidget(browse_btn)
        template_btn = QPushButton('下载模板')
        template_btn.clicked.connect(lambda: self._download_template(
            'addr' if self.batch_mode_combo.currentText() == '地址转坐标' else 'coord'))
        fg.addWidget(template_btn)
        file_grp.setLayout(fg)
        layout.addWidget(file_grp)

        self.preview_table = QTableWidget()
        self.preview_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        layout.addWidget(self.preview_table, 1)

        set_grp = QGroupBox('转换设置')
        sg = QGridLayout()
        sg.setHorizontalSpacing(12)
        sg.setVerticalSpacing(8)
        sg.addWidget(QLabel('转换模式：'), 0, 0)
        self.batch_mode_combo = QComboBox()
        self.batch_mode_combo.addItems(['地址转坐标', '坐标转地址'])
        self.batch_mode_combo.currentTextChanged.connect(self._on_batch_mode_changed)
        sg.addWidget(self.batch_mode_combo, 0, 1)
        sg.addWidget(QLabel('地址列：'), 0, 2)
        self.addr_col_combo = QComboBox()
        sg.addWidget(self.addr_col_combo, 0, 3)
        sg.addWidget(QLabel('经度列：'), 1, 0)
        self.b_lon_combo = QComboBox()
        sg.addWidget(self.b_lon_combo, 1, 1)
        sg.addWidget(QLabel('纬度列：'), 1, 2)
        self.b_lat_combo = QComboBox()
        sg.addWidget(self.b_lat_combo, 1, 3)
        sg.addWidget(QLabel('坐标系列：'), 2, 0)
        self.b_crs_combo = QComboBox()
        sg.addWidget(self.b_crs_combo, 2, 1)
        sg.addWidget(QLabel('默认城市：'), 2, 2)
        self.b_city_edit = QLineEdit()
        sg.addWidget(self.b_city_edit, 2, 3)
        set_grp.setLayout(sg)
        layout.addWidget(set_grp)

        top = QHBoxLayout()
        top.addStretch()
        self.batch_start_btn = QPushButton('▶ 开始批量转换')
        self.batch_start_btn.clicked.connect(self._start_batch)
        top.addWidget(self.batch_start_btn)
        self.batch_cancel_btn = QPushButton('取消')
        self.batch_cancel_btn.setEnabled(False)
        self.batch_cancel_btn.clicked.connect(self._cancel_batch)
        top.addWidget(self.batch_cancel_btn)
        self.batch_open_btn = QPushButton('打开文件夹')
        self.batch_open_btn.clicked.connect(self._open_batch_folder)
        top.addWidget(self.batch_open_btn)
        top.addStretch()
        layout.addLayout(top)
        self.batch_progress = QProgressBar()
        self.batch_progress.setVisible(False)
        layout.addWidget(self.batch_progress)

        self.batch_log = QTextEdit()
        self.batch_log.setReadOnly(True)
        self.batch_log.setFixedHeight(120)
        layout.addWidget(self.batch_log)

        self.export_btn = QPushButton('导出结果')
        self.export_btn.setVisible(False)

        self.batch_input_path = ''
        self.batch_input_count = 0
        self.df_headers = []
        self.batch_results = []
        self.batch_failures = []
        self.batch_output_path = ''
        self._on_batch_mode_changed()
        return page

    def _on_batch_mode_changed(self):
        is_addr = self.batch_mode_combo.currentText() == '地址转坐标'
        self.addr_col_combo.setEnabled(is_addr)
        self.b_lon_combo.setEnabled(not is_addr)
        self.b_lat_combo.setEnabled(not is_addr)
        self.b_crs_combo.setEnabled(not is_addr)

    def _browse(self):
        path, _ = QFileDialog.getOpenFileName(
            self, '选择文件', '', '表格文件 (*.xlsx *.xls *.csv);;Excel (*.xlsx *.xls);;CSV (*.csv)')
        if path:
            self.file_edit.setText(path)
            self._load_file()

    def _load_file(self):
        path = self.file_edit.text().strip()
        if not path:
            QMessageBox.warning(self, '提示', '请先选择文件。')
            return
        ext = os.path.splitext(path)[1].lower()
        try:
            info = inspect_table(path)
        except Exception as e:
            QMessageBox.critical(self, '读取失败', str(e))
            return
        self.df_headers = [str(h).strip() for h in info['headers']]
        self.batch_input_path = path
        self.batch_input_count = int(info['count'])
        # 填充列下拉
        cols = self.df_headers
        for combo in (self.addr_col_combo, self.b_lon_combo, self.b_lat_combo, self.b_crs_combo):
            combo.clear()
            combo.addItems(cols)
        for i, col in enumerate(cols):
            cl = col.lower()
            if '地址' in col or 'addr' in cl:
                self.addr_col_combo.setCurrentIndex(i)
            if '经度' in col or 'lng' in cl or 'lon' in cl:
                self.b_lon_combo.setCurrentIndex(i)
            if '纬度' in col or 'lat' in cl:
                self.b_lat_combo.setCurrentIndex(i)
            if '坐标系' in col or 'crs' in cl:
                self.b_crs_combo.setCurrentIndex(i)
        # 预览前 5 行
        preview_rows = info['preview']
        n = len(preview_rows)
        self.preview_table.setColumnCount(len(cols))
        self.preview_table.setHorizontalHeaderLabels(cols)
        self.preview_table.setRowCount(n)
        for r in range(n):
            for c in range(len(cols)):
                val = preview_rows[r][c] if c < len(preview_rows[r]) else ''
                self.preview_table.setItem(r, c, QTableWidgetItem('' if val is None else str(val)))
        self.b_city_edit.setText(load_config().get('defaults', {}).get('city', ''))
        self.batch_log.append(f'已识别 {self.batch_input_count:,} 条记录；仅预览前 {n} 行')

    def _read_csv(self, path):
        rows = None
        last = None
        for enc in CSV_ENCODINGS:
            try:
                with open(path, 'r', encoding=enc, newline='') as f:
                    rows = list(csv.reader(f))
                break
            except UnicodeDecodeError as e:
                last = e
        if rows is None:
            raise UnicodeDecodeError('无法解码 CSV') from last
        return rows

    def _read_excel(self, path):
        wb = load_workbook(path, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]

    def _download_template(self, kind):
        path, _ = QFileDialog.getSaveFileName(
            self, '保存模板', '地址转坐标模板.xlsx' if kind == 'addr' else '坐标转地址模板.xlsx',
            'Excel (*.xlsx);;CSV (*.csv)')
        if not path:
            return
        header = ['地址'] if kind == 'addr' else ['经度', '纬度', '坐标系']
        sample = [['南宁市民族大道1号'], ['南宁市朝阳路2号']] if kind == 'addr' \
            else [[108.3667, 22.8167, 'GCJ-02'], [108.3812, 22.8214, 'GCJ-02']]
        try:
            if path.lower().endswith('.csv'):
                with open(path, 'w', newline='', encoding='utf-8-sig') as f:
                    w = csv.writer(f)
                    w.writerow(header)
                    w.writerows(sample)
            else:
                wb = Workbook()
                ws = wb.active
                ws.append(header)
                for s in sample:
                    ws.append(s)
                wb.save(path)
            QMessageBox.information(self, '成功', f'模板已保存：{path}')
        except Exception as e:
            QMessageBox.critical(self, '失败', str(e))

    def _start_batch(self):
        if not self.batch_input_path:
            QMessageBox.warning(self, '无数据', '请先选择并打开数据文件。')
            return
        mode = 'addr2coord' if self.batch_mode_combo.currentText() == '地址转坐标' else 'coord2addr'
        cfg = load_config()
        defaults = cfg.get('defaults', {})
        province = defaults.get('province', '')
        city = self.b_city_edit.text().strip() or defaults.get('city', '')
        addr_col = self.addr_col_combo.currentText()
        lon_col = self.b_lon_combo.currentText()
        lat_col = self.b_lat_combo.currentText()
        crs_col = self.b_crs_combo.currentText()
        output_path, _ = QFileDialog.getSaveFileName(
            self, '保存批量处理结果',
            f'{self.batch_mode_combo.currentText()}_{datetime.now().strftime("%Y%m%d_%H%M%S")}_V2.0.xlsx',
            'Excel (*.xlsx);;CSV (*.csv)')
        if not output_path:
            return
        self.batch_start_btn.setEnabled(False)
        self.batch_cancel_btn.setEnabled(True)
        self.export_btn.setEnabled(False)
        self.batch_progress.setVisible(True)
        self.batch_progress.setValue(0)
        self.batch_log.append('开始批量转换…')
        self.worker = BatchWorker(mode, self.batch_input_path, output_path,
                                  addr_col, lon_col, lat_col, crs_col, city, province)
        self.worker.progress.connect(lambda i, t: self.batch_progress.setValue(int(i / t * 100)))
        self.worker.log.connect(self.batch_log.append)
        self.worker.finished_ok.connect(self._on_batch_finished)
        self.worker.error.connect(self._on_batch_error)
        self.worker.start()

    def _on_batch_finished(self, summary):
        self.batch_results = summary['preview']
        self.batch_failures = summary['failed']
        self.batch_output_path = summary['output_path']
        self.batch_progress.setVisible(False)
        self.batch_start_btn.setEnabled(True)
        self.batch_cancel_btn.setEnabled(False)
        self.export_btn.setEnabled(bool(summary['total']))
        self.batch_log.append(f'完成：成功 {summary["success"]:,} 条，失败 {summary["failed"]:,} 条')

    def _on_batch_error(self, message):
        self.batch_progress.setVisible(False)
        self.batch_start_btn.setEnabled(True)
        self.batch_cancel_btn.setEnabled(False)
        self.export_btn.setEnabled(False)
        QMessageBox.critical(self, '处理失败', message)

    def _export_batch(self):
        if not self.batch_results:
            QMessageBox.warning(self, '无数据', '没有可导出的结果。')
            return
        path, _ = QFileDialog.getSaveFileName(
            self, '另存结果', f'地址坐标互转结果_导出{os.path.splitext(self.batch_output_path)[1]}',
            'Excel (*.xlsx)' if self.batch_output_path.lower().endswith('.xlsx') else 'CSV (*.csv)')
        if not path:
            return
        try:
            if os.path.abspath(path) == os.path.abspath(self.batch_output_path):
                QMessageBox.information(self, '提示', '结果已保存在该路径，无需重复复制。')
                return
            shutil.copy2(self.batch_output_path, path)
            QMessageBox.information(self, '成功', f'结果已导出：{path}')
        except Exception as e:
            QMessageBox.critical(self, '导出失败', str(e))

    def _cancel_batch(self):
        if getattr(self, 'worker', None) and self.worker.isRunning():
            self.worker.stop()
            self.batch_log.append('正在取消；已完成结果将保留。')

    def _open_batch_folder(self):
        path = os.path.dirname(self.batch_output_path) if self.batch_output_path else os.path.join(os.getcwd(), 'output')
        os.makedirs(path, exist_ok=True)
        subprocess.Popen(['explorer', path])
