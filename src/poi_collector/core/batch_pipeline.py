# -*- coding: utf-8 -*-
"""Bounded-memory CSV/XLSX inspection and coordinate conversion pipelines."""
from __future__ import annotations

import csv
import os
from pathlib import Path

from core.coords import transform_between

CSV_ENCODINGS = ('utf-8-sig', 'gb18030', 'gbk', 'utf-8')


def detect_csv_encoding(path):
    for encoding in CSV_ENCODINGS:
        try:
            with open(path, 'r', encoding=encoding, newline='') as stream:
                stream.read(65536)
            return encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeError(f'无法识别 CSV 编码：{path}')


def inspect_table(path, preview_rows=5):
    extension = Path(path).suffix.lower()
    if extension == '.csv':
        encoding = detect_csv_encoding(path)
        with open(path, 'r', encoding=encoding, newline='') as stream:
            reader = csv.reader(stream)
            headers = next(reader, [])
            preview = []
            count = 0
            for row in reader:
                count += 1
                if len(preview) < preview_rows:
                    preview.append(row)
        return {'headers': headers, 'preview': preview, 'count': count, 'encoding': encoding}

    if extension == '.xls':
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError('读取 .xls 需要 xlrd 组件，请重新安装完整版程序') from exc
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            sheet = book.sheet_by_index(0)
            headers = sheet.row_values(0) if sheet.nrows else []
            preview = [sheet.row_values(i) for i in range(1, min(sheet.nrows, preview_rows + 1))]
            return {'headers': headers, 'preview': preview, 'count': max(sheet.nrows - 1, 0),
                    'encoding': None}
        finally:
            book.release_resources()
    if extension != '.xlsx':
        raise ValueError(f'不支持的表格格式：{extension}；请选择 xlsx、xls 或 csv')
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = list(next(iterator, ()))
        preview = [list(row) for _, row in zip(range(preview_rows), iterator)]
        return {'headers': headers, 'preview': preview, 'count': max(sheet.max_row - 1, 0),
                'encoding': None}
    finally:
        workbook.close()


def iter_rows(path, headers, encoding=None):
    extension = Path(path).suffix.lower()
    if extension == '.csv':
        with open(path, 'r', encoding=encoding or detect_csv_encoding(path), newline='') as stream:
            yield from csv.DictReader(stream)
        return
    if extension == '.xls':
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            sheet = book.sheet_by_index(0)
            for row_index in range(1, sheet.nrows):
                values = sheet.row_values(row_index)
                yield {header: values[index] if index < len(values) else ''
                       for index, header in enumerate(headers)}
        finally:
            book.release_resources()
        return
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        next(iterator, None)
        for values in iterator:
            yield {header: values[index] if index < len(values) else ''
                   for index, header in enumerate(headers)}
    finally:
        workbook.close()


def stream_coordinate_conversion(input_path, output_path, src_cs, dst_cs,
                                 name_col='点位名称', lon_col='源经度', lat_col='源纬度',
                                 progress_cb=None, cancel_cb=None, preview_limit=500):
    info = inspect_table(input_path, preview_rows=0)
    headers = [str(item).strip() for item in info['headers']]
    for required in (lon_col, lat_col):
        if required not in headers:
            raise ValueError(f'缺少必需列：{required}')
    total = int(info['count'])
    preview = []
    success = failed = 0
    output_headers = [name_col, lon_col, lat_col, '目标经度', '目标纬度', '状态']
    extension = Path(output_path).suffix.lower()

    csv_stream = None
    workbook = None
    sheet = None
    rows_in_sheet = 0
    sheet_index = 0
    if extension == '.csv':
        csv_stream = open(output_path, 'w', encoding='utf-8-sig', newline='')
        writer = csv.writer(csv_stream)
        writer.writerow(output_headers)
    else:
        from openpyxl import Workbook
        workbook = Workbook(write_only=True)
        writer = None

    try:
        for index, row in enumerate(iter_rows(input_path, headers, info['encoding']), 1):
            if cancel_cb and cancel_cb():
                break
            name = row.get(name_col, '') if name_col in headers else ''
            raw_lon, raw_lat = row.get(lon_col, ''), row.get(lat_col, '')
            try:
                lon, lat = float(raw_lon), float(raw_lat)
                out_lon, out_lat = transform_between(lon, lat, src_cs, dst_cs)
                output = [name, lon, lat, round(out_lon, 8), round(out_lat, 8), '成功']
                success += 1
            except Exception as exc:
                output = [name, raw_lon, raw_lat, '', '', f'失败：{exc}']
                failed += 1
            if csv_stream:
                writer.writerow(output)
            else:
                if sheet is None or rows_in_sheet >= 1_000_000:
                    sheet_index += 1
                    sheet = workbook.create_sheet(f'转换结果{sheet_index:02d}')
                    sheet.append(output_headers)
                    rows_in_sheet = 0
                sheet.append(output)
                rows_in_sheet += 1
            if len(preview) < preview_limit:
                preview.append(output)
            if progress_cb and (index == total or index % 10_000 == 0):
                progress_cb(index, total)
    finally:
        if csv_stream:
            csv_stream.close()
        elif workbook:
            workbook.save(output_path)

    return {'success': success, 'failed': failed, 'total': success + failed,
            'preview': preview, 'output_path': str(output_path)}


def stream_address_conversion(input_path, output_path, mode, addr_col='', lon_col='',
                              lat_col='', crs_col='', city='', province='',
                              progress_cb=None, log_cb=None, cancel_cb=None,
                              preview_limit=200):
    """Stream address/geocode API work directly to disk with bounded memory."""
    from core import geo_service
    from core.coords import transform_between
    from core.utils import parse_coordinate_str, validate_coordinate

    info = inspect_table(input_path, preview_rows=0)
    headers = [str(item).strip() for item in info['headers']]
    total = int(info['count'])
    if mode == 'addr2coord':
        output_headers = ['输入地址', '标准地址', 'WGS84经度', 'WGS84纬度',
                          'GCJ02经度', 'GCJ02纬度', 'BD09经度', 'BD09纬度', '来源', '状态']
    else:
        output_headers = ['输入经度', '输入纬度', '输入坐标系', '标准地址',
                          '省', '市', '区县', '乡镇街道', '门牌号', '来源', '状态']

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    csv_stream = None
    workbook = None
    sheet = None
    rows_in_sheet = 0
    sheet_index = 0
    if Path(output_path).suffix.lower() == '.csv':
        csv_stream = open(output_path, 'w', encoding='utf-8-sig', newline='')
        writer = csv.DictWriter(csv_stream, fieldnames=output_headers)
        writer.writeheader()
    else:
        from openpyxl import Workbook
        workbook = Workbook(write_only=True)
        writer = None

    preview = []
    success = failed = 0
    preferred = 'amap'
    try:
        from core.config import load_config
        preferred = load_config().get('defaults', {}).get('preferred_provider', 'amap')
        for index, row in enumerate(iter_rows(input_path, headers, info['encoding']), 1):
            if cancel_cb and cancel_cb():
                break
            try:
                if mode == 'addr2coord':
                    address = str(row.get(addr_col, '') or '').strip()
                    result = geo_service.geocode_with_fallback(address, city, province, preferred) if address else None
                    if not result:
                        raise ValueError('地址为空或所有可用图源均未命中')
                    wgs, gcj, bd = result['wgs84'], result['gcj02'], result['bd09']
                    output = {
                        '输入地址': address, '标准地址': result.get('formatted', ''),
                        'WGS84经度': round(wgs[0], 6), 'WGS84纬度': round(wgs[1], 6),
                        'GCJ02经度': round(gcj[0], 6), 'GCJ02纬度': round(gcj[1], 6),
                        'BD09经度': round(bd[0], 6), 'BD09纬度': round(bd[1], 6),
                        '来源': result.get('provider', ''), '状态': '成功',
                    }
                else:
                    raw_lon, raw_lat = row.get(lon_col, ''), row.get(lat_col, '')
                    input_crs = str(row.get(crs_col, '') or 'WGS84').strip().upper() if crs_col else 'WGS84'
                    lon, lat = parse_coordinate_str(raw_lon), parse_coordinate_str(raw_lat)
                    if lon is None or lat is None or not validate_coordinate(lon, lat):
                        raise ValueError('坐标格式无效')
                    wlon, wlat = transform_between(lon, lat, input_crs, 'WGS84')
                    result = geo_service.reverse_full(wlon, wlat, preferred)
                    if not result:
                        raise ValueError('所有可用图源均未命中')
                    output = {
                        '输入经度': raw_lon, '输入纬度': raw_lat, '输入坐标系': input_crs,
                        '标准地址': result.get('formatted', ''), '省': result.get('province', ''),
                        '市': result.get('city', ''), '区县': result.get('district', ''),
                        '乡镇街道': ' '.join(filter(None, (result.get('township', ''), result.get('street', '')))),
                        '门牌号': result.get('number', ''), '来源': result.get('provider', ''), '状态': '成功',
                    }
                success += 1
            except Exception as exc:
                failed += 1
                output = {key: '' for key in output_headers}
                if mode == 'addr2coord':
                    output['输入地址'] = str(row.get(addr_col, '') or '')
                else:
                    output['输入经度'], output['输入纬度'] = row.get(lon_col, ''), row.get(lat_col, '')
                output['状态'] = f'失败：{exc}'
                if log_cb:
                    log_cb(f'[行 {index + 1}] {output["状态"]}')

            if csv_stream:
                writer.writerow(output)
            else:
                if sheet is None or rows_in_sheet >= 1_000_000:
                    sheet_index += 1
                    sheet = workbook.create_sheet(f'处理结果{sheet_index:02d}')
                    sheet.append(output_headers)
                    rows_in_sheet = 0
                sheet.append([output.get(key, '') for key in output_headers])
                rows_in_sheet += 1
            if len(preview) < preview_limit:
                preview.append(output)
            if progress_cb and (index == total or index % 100 == 0):
                progress_cb(index, total)
            if log_cb and index % 1000 == 0:
                log_cb(f'[进度] {index:,}/{total:,}，成功 {success:,}，失败 {failed:,}')
    finally:
        if csv_stream:
            csv_stream.close()
        elif workbook:
            workbook.save(output_path)
    return {'success': success, 'failed': failed, 'total': success + failed,
            'preview': preview, 'output_path': str(output_path)}
