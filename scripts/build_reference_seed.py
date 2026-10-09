# -*- coding: utf-8 -*-
"""把用户核验的六 Sheet 区划工作簿转换为程序离线缓存。"""
from __future__ import annotations

import csv
import json
import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DIVISION_OUTPUT = PROJECT_ROOT / 'data' / 'divisions' / '450000000.json'
REFERENCE_DIVISION_OUTPUT = PROJECT_ROOT / 'data' / 'divisions' / '450000000_reference_v1.json'
VILLAGE_OUTPUT = PROJECT_ROOT / 'data' / 'villages' / '450000000.csv'


def text_code(value, length):
    if value in (None, ''):
        return ''
    if isinstance(value, (int, float)):
        value = str(int(value))
    else:
        value = str(value).strip()
        if 'e+' in value.lower():
            value = str(int(float(value)))
    return value[:length].ljust(length, '0')


def rows_by_header(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or '').strip() for value in next(rows)]
    for values in rows:
        yield {headers[index]: values[index] for index in range(min(len(headers), len(values)))}


def provider_from_chain(chain):
    chain = str(chain or '')
    if '天地图' in chain:
        return 'tianditu'
    if '高德' in chain:
        return 'amap'
    if '百度' in chain:
        return 'baidu'
    if '腾讯' in chain:
        return 'tencent'
    return '本地核验缓存'


def node(row, level, name_col, code_col, parent_code):
    lengths = {1: 9, 2: 9, 3: 9, 4: 9}
    return {
        'code': text_code(row.get(code_col), lengths[level]),
        'name': str(row.get(name_col) or '').strip(),
        'level': level,
        'parent': parent_code,
        '政府驻地地址': str(row.get('政府驻地地址') or '').strip(),
        '经度(WGS84)': row.get('政府驻地经度(WGS84)') or '',
        '纬度(WGS84)': row.get('政府驻地纬度(WGS84)') or '',
        '匹配源': provider_from_chain(row.get('坐标纠偏链')),
        '坐标纠偏链': str(row.get('坐标纠偏链') or ''),
        '数据来源': str(row.get('数据来源') or ''),
    }


def main(source: Path):
    workbook = load_workbook(source, read_only=True, data_only=True)
    nodes = [{
        'code': '450000000', 'name': '广西壮族自治区', 'level': 1, 'parent': '',
        '匹配源': '本地核验缓存', '数据来源': '用户核验工作簿 V1.0',
    }]
    cities = {}
    for row in rows_by_header(workbook['地市表']):
        item = node(row, 2, '地市名称', '地市编码', '450000000')
        nodes.append(item)
        cities[item['code'][:4]] = item['code']
    districts = {}
    for row in rows_by_header(workbook['区县表']):
        city_prefix = text_code(row.get('地市编码'), 4)[:4]
        item = node(row, 3, '区县名称', '区划编码', cities.get(city_prefix, city_prefix.ljust(9, '0')))
        nodes.append(item)
        districts[item['code'][:6]] = item['code']
    towns = {}
    for row in rows_by_header(workbook['乡镇表']):
        district_prefix = text_code(row.get('区县编码'), 6)[:6]
        item = node(
            row, 4, '乡镇名称', '乡镇编码',
            districts.get(district_prefix, district_prefix.ljust(9, '0')),
        )
        nodes.append(item)
        towns[item['code'][:9]] = item['code']

    DIVISION_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    division_payload = json.dumps({
        'meta': {
            'complete': True, 'max_level': 4,
            'updated_at': datetime.now().isoformat(timespec='seconds'),
            'source': '广西行政区划数据_V1.0.xlsx（用户核验，Workbuddy生成）',
            'counts': {'city': len(cities), 'district': len(districts), 'town': len(towns)},
        },
        'nodes': nodes,
    }, ensure_ascii=False, indent=1)
    DIVISION_OUTPUT.write_text(division_payload, encoding='utf-8')
    REFERENCE_DIVISION_OUTPUT.write_text(division_payload, encoding='utf-8')

    village_headers = [
        '编码', '名称', '乡镇编码', '地市名称', '地市编码', '区县名称', '区县编码',
        '乡镇名称', '村级分类', '经度(WGS84)', '纬度(WGS84)', '匹配源',
        '数据来源', '数据版本',
    ]
    VILLAGE_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    village_count = 0
    with VILLAGE_OUTPUT.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=village_headers)
        writer.writeheader()
        for row in rows_by_header(workbook['行政村详细清单']):
            code = text_code(row.get('行政村编码'), 12)
            town_code = text_code(row.get('乡镇编码'), 9)
            data_source = str(row.get('数据来源') or '')
            writer.writerow({
                '编码': code, '名称': row.get('行政村名称') or '',
                '乡镇编码': town_code, '地市名称': row.get('地市名称') or '',
                '地市编码': text_code(row.get('地市编码'), 6),
                '区县名称': row.get('区县名称') or '',
                '区县编码': text_code(row.get('区县编码'), 6),
                '乡镇名称': row.get('乡镇名称') or '',
                '村级分类': row.get('村级分类') or '',
                '经度(WGS84)': row.get('经度(WGS84)') or '',
                '纬度(WGS84)': row.get('纬度(WGS84)') or '',
                '匹配源': provider_from_chain(data_source),
                '数据来源': data_source,
                '数据版本': '广西行政区划数据_V1.0（用户核验）',
            })
            village_count += 1
    print(json.dumps({
        'division_nodes': len(nodes), 'villages': village_count,
        'division_output': str(DIVISION_OUTPUT),
        'reference_division_output': str(REFERENCE_DIVISION_OUTPUT),
        'village_output': str(VILLAGE_OUTPUT),
    }, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='从核验工作簿生成离线区划种子数据')
    parser.add_argument('workbook', type=Path, help='输入工作簿路径（XLSX）')
    args = parser.parse_args()
    main(args.workbook.resolve())
