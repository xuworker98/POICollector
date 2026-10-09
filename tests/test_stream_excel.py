from openpyxl import load_workbook

from core.excel_out import write_poi_store_excel
from core.store import PoiStore


def test_stream_poi_excel_has_six_logical_sections(tmp_path):
    store = PoiStore(tmp_path / 'poi.sqlite3')
    store.insert_pois([{
        'key': 'a', 'name': '学校A', 'type': '学校', 'district': '青秀区',
        'lon': 108.3, 'lat': 22.8, 'source': 'amap', 'confidence': '中',
    }])
    store.close()
    output = tmp_path / 'poi.xlsx'
    write_poi_store_excel(output, tmp_path / 'poi.sqlite3')
    workbook = load_workbook(output, read_only=False)
    try:
        assert workbook.sheetnames == [
            'POI明细01', '汇总表', '来源统计', '置信度统计', '校验表', '运行信息',
        ]
        headers = [cell.value for cell in next(workbook['POI明细01'].iter_rows())]
        assert headers == [
            '省份', '地市', '区县', 'POI名称', 'POI分类', '经度(WGS84)',
            '纬度(WGS84)', 'POI地址', '数据来源', '置信度', '行政区码', '分类码',
        ]
        assert workbook['POI明细01'].freeze_panes == 'A2'
        assert workbook['POI明细01'].auto_filter.ref == 'A1:L2'
        assert workbook['汇总表'].freeze_panes == 'A2'
        assert workbook['汇总表'].auto_filter.ref == 'A1:B3'
    finally:
        workbook.close()
